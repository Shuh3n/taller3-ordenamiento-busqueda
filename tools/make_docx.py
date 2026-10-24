"""Builds a Word/Google-Docs version (report/Informe_Taller3.docx) from report/overleaf/main.tex.

The LaTeX source stays the single source of truth: this script flattens it (inlines the generated tables,
resolves numbering and cross-references, expands the figure macro), runs pandoc, and then post-processes the
result with python-docx (table borders, centered figures and cover, page breaks, page numbers).
Requires pandoc on PATH and python-docx. Usage: python tools/make_docx.py
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "report" / "overleaf"
OUT = ROOT / "report" / "Informe_Taller3.docx"
BREAK = "@@PAGEBREAK@@"

COVER = [
    "Universidad del Quindío", "@@GAP@@", "Taller 3", "Ordenamiento y Búsqueda",
    "Comparación experimental de algoritmos de ordenamiento y búsqueda en Java y Python", "@@GAP@@",
    "Santiago Orozco Zuluaga", "Jhony Alexander Perea Perea", "@@GAP@@",
    "Curso: Análisis de Algoritmos", "Docente: Carlos Andrés Flórez Villarraga",
    "Repositorio: https://github.com/Shuh3n/taller3-ordenamiento-busqueda", "@@GAP@@", "Octubre de 2026",
]


def group(s, i):
    """s[i] == '{': return (content, index after the matching '}')."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{" and s[j - 1] != "\\":
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces")


def pop_cmd(body, name):
    """Remove \\name{...} from body; return (content or None, new body)."""
    k = body.find("\\" + name + "{")
    if k < 0:
        return None, body
    content, end = group(body, k + len(name) + 1)
    return content, body[:k] + body[end:]


def flatten(tex):
    body = tex.split("\\begin{document}", 1)[1].split("\\end{document}", 1)[0]
    body = re.sub(r"\\input\{([^}]*)\}", lambda m: (SRC / m.group(1)).read_text(encoding="utf-8"), body)
    body = re.sub(r"\\begin\{titlepage\}.*?\\end\{titlepage\}", "@@COVER@@", body, flags=re.S)
    body = body.replace("\\tableofcontents\n\\clearpage", "").replace("\\clearpage", BREAK)

    labels = {"ref": {}, "cite": {}}

    # sections ------------------------------------------------------------------------------
    state = {"sec": 0, "sub": 0, "app": False}

    def sec(m):
        if m.group(1):
            state["app"], state["sec"], state["sub"] = True, 0, 0
            return ""
        kind, title, label = m.group(2), m.group(3), m.group(5)
        if kind == "section":
            state["sec"] += 1
            state["sub"] = 0
            num = "A" if state["app"] else str(state["sec"])
            heading = f"Anexo A. {title}" if state["app"] else f"{num}. {title}"
            cmd = "section"
        else:
            state["sub"] += 1
            num = f"{'A' if state['app'] else state['sec']}.{state['sub']}"
            heading, cmd = f"{num}. {title}", "subsection"
        if label:
            labels["ref"][label] = num
        return f"\\{cmd}*{{{heading}}}"

    body = re.sub(r"(\\appendix)|\\(section|subsection)\{((?:[^{}]|\{[^{}]*\})*)\}(\s*\\label\{([^}]*)\})?",
                  sec, body)

    # tables, figures, equations ---------------------------------------------------------------
    counters = {"t": 0, "f": 0, "e": 0}

    def table(m):
        counters["t"] += 1
        n = counters["t"]
        inner = m.group(2)
        cap, inner = pop_cmd(inner, "caption")
        lab_in_cap = re.search(r"\\label\{([^}]*)\}", inner)
        if lab_in_cap:
            labels["ref"][lab_in_cap.group(1)] = str(n)
        inner = re.sub(r"\\label\{[^}]*\}|\\centering", "", inner)
        return f"\n\n\\textbf{{Tabla {n}.}} {cap}\n\n{inner}\n\n"

    body = re.sub(r"\\begin\{table\}(\[[^\]]*\])?(.*?)\\end\{table\}", table, body, flags=re.S)

    out, i = [], 0
    while True:
        k = body.find("\\figpair{", i)
        if k < 0:
            out.append(body[i:])
            break
        out.append(body[i:k])
        j, args = k + len("\\figpair"), []
        for _ in range(4):
            a, j = group(body, j)
            args.append(a)
        name, width, cap, label = args
        counters["f"] += 1
        n = counters["f"]
        labels["ref"][label] = str(n)
        for machine, letter in (("A", "a"), ("B", "b")):
            out.append(f"\n\n\\includegraphics[width={width}\\linewidth]{{figs/{machine}_{name}.png}}\n\n"
                       f"\\textit{{Figura {n}{letter}. {cap} (Equipo {machine}).}}\n\n")
        i = j
    body = "".join(out)

    def eq(m):
        counters["e"] += 1
        n = counters["e"]
        inner = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", inner)
        if lab:
            labels["ref"][lab.group(1)] = str(n)
        inner = re.sub(r"\\label\{[^}]*\}", "", inner)
        return f"\\begin{{equation}}{inner}\\qquad ({n})\\end{{equation}}"

    body = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}", eq, body, flags=re.S)

    # bibliography ------------------------------------------------------------------------------
    def bib(m):
        parts = re.split(r"\\bibitem\{([^}]*)\}", m.group(1))
        res = ["\n\n\\section*{Referencias}\n\n"]
        for k in range(1, len(parts), 2):
            labels["cite"][parts[k]] = str(k // 2 + 1)
            res.append(f"[{k // 2 + 1}] {parts[k + 1].strip()}\n\n")
        return "".join(res)

    body = re.sub(r"\\begin\{thebibliography\}\{[^}]*\}(.*?)\\end\{thebibliography\}", bib, body, flags=re.S)

    # references and macros --------------------------------------------------------------------
    body = re.sub(r"\\eqref\{([^}]*)\}", lambda m: f"({labels['ref'][m.group(1)]})", body)
    body = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels["ref"][m.group(1)], body)
    body = re.sub(r"\\cite\{([^}]*)\}", lambda m: f"[{labels['cite'][m.group(1)]}]", body)
    body = re.sub(r"\\label\{[^}]*\}", "", body)
    # simple inline math -> plain Unicode text (keeps Word/Docs from filling up with tiny equation objects)
    symbols = {"\\times": "×", "\\sim": "~", "\\pm": "±", "\\to": "→", "\\mu": "μ", "\\Rightarrow": "⇒",
               "\\approx": "≈", ">": ">"}
    for cmd, sym in symbols.items():
        body = re.sub(r"\$" + re.escape(cmd) + r"\$", sym, body)
    body = re.sub(r"\$(\\times|\\approx)\s*([0-9]+(?:\{,\}[0-9]+)?)\$",
                  lambda m: symbols[m.group(1)] + m.group(2).replace("{,}", ","), body)
    body = re.sub(r"\$([a-zA-Z])\$", r"\\emph{\1}", body)
    body = body.replace("\\,", " ")
    body = re.sub(r"\\paragraph\{([^}]*)\}", r"\\textbf{\1}", body)
    body = re.sub(r"\\pendiente\{((?:[^{}]|\{[^{}]*\})*)\}", r"\\textbf{[\1]}", body)
    body = re.sub(r"\\code\{", r"\\texttt{", body)
    body = body.replace("\\-", "").replace("\\textquestiondown", "¿").replace("\\LaTeX{}", "LaTeX")
    body = re.sub(r"\\Needspace\{[^}]*\}|\\resizebox\{\\linewidth\}\{!\}|\\medskip|\\smallskip", "", body)
    body = re.sub(r"\\(small|footnotesize|scriptsize|tiny)\b", "", body)
    body = re.sub(r"\{\s*(\\begin\{tabular\})", r"\1", body)
    body = body.replace("\\end{tabular}}", "\\end{tabular}")
    body = body.replace("<<", "«").replace(">>", "»")

    cover = "\n\n".join(f"\\textbf{{{line}}}" if line == "Taller 3" else line for line in COVER)
    return body.replace("@@COVER@@", cover + f"\n\n{BREAK}\n\n")


# ---------------------------------------------------------------------------------- post-process
def set_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


def shade(cell, color="E8EEF4"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def postprocess(path):
    doc = Document(str(path))
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.4)
    page_number_footer(doc.sections[0])
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    for t in doc.tables:
        set_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                if ri == 0:
                    shade(cell)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        if ri == 0:
                            r.bold = True

    cover_done = False
    for idx, p in enumerate(list(doc.paragraphs)):
        text = p.text.strip()
        if text == BREAK:
            for r in p.runs:
                r.text = ""
            p.add_run().add_break(WD_BREAK.PAGE)
            cover_done = True
            continue
        if not cover_done:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if text == "@@GAP@@":
                for r in p.runs:
                    r.text = ""
                p.paragraph_format.space_after = Pt(36)
            elif text == "Taller 3":
                for r in p.runs:
                    r.font.size = Pt(30)
                    r.bold = True
            elif text == "Ordenamiento y Búsqueda":
                for r in p.runs:
                    r.font.size = Pt(22)
                    r.bold = True
            continue
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
        elif text.startswith("Figura ") and re.match(r"Figura \d+[ab]\.", text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)
        elif re.match(r"Tabla \d+\.", text):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
    doc.save(str(path))


def main():
    flat = flatten((SRC / "main.tex").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as tmp:
        tex = Path(tmp) / "flat.tex"
        tex.write_text(flat, encoding="utf-8")
        cmd = ["pandoc", str(tex), "-f", "latex", "-t", "docx", "-o", str(OUT), "--resource-path", str(SRC)]
        res = subprocess.run(cmd, capture_output=True, text=True)
        if res.returncode:
            sys.exit(res.stderr)
        if res.stderr.strip():
            print(res.stderr.strip()[:2000])
    postprocess(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
